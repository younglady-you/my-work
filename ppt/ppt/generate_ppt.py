# -*- coding: utf-8 -*-
"""
generate_ppt.py
生成基础 PPTX（根据仓库中 ppt/试用期工作总结_SSC_薪酬核算专员.txt 的分段生成幻灯片）
依赖: python-pptx
安装: pip install python-pptx
运行: python3 ppt/generate_ppt.py
输出: ppt/试用期工作总结_SSC_薪酬核算专员.pptx
"""
from pptx import Presentation
from pptx.util import Pt, Inches
from pptx.dml.color import RGBColor
import os
import sys

# 输入/输出
INPUT_TXT = os.path.join('ppt', '试用期工作总结_SSC_薪酬核算专员.txt')
OUTPUT_PPTX = os.path.join('ppt', '试用期工作总结_SSC_薪酬核算专员.pptx')

# 字体/颜色设置（在 runner 上未必有微软雅黑，打开后可在 PowerPoint 调整）
TITLE_FONT = 'Microsoft YaHei'  # 如果不可用会使用默认字体
BODY_FONT = 'Microsoft YaHei'
TITLE_COLOR = RGBColor(31, 78, 121)   # #1F4E79 深蓝
ACCENT_COLOR = RGBColor(244, 162, 97) # #F4A261 橙

def read_slides_from_txt(path):
    if not os.path.exists(path):
        print(f"Input file not found: {path}")
        return []
    with open(path, 'r', encoding='utf-8') as f:
        txt = f.read()
    # 按 --- 幻灯片 分段（保持兼容之前的文本格式）
    parts = [p.strip() for p in txt.split('--- 幻灯片') if p.strip()]
    slides = []
    for p in parts:
        # p 可能以 " 1 封面 ---\n内容" 的形式存在，我们尝试提取标题
        # 找第一个 '---' 分隔（若存在）
        if ' ---' in p:
            idx = p.find(' ---')
            title = p[:idx].strip()
            body = p[idx+4:].strip()
        else:
            # 否则以首行作为标题
            lines = p.splitlines()
            title = lines[0].strip() if lines else ''
            body = '\n'.join(lines[1:]).strip() if len(lines) > 1 else ''
        slides.append((title, body))
    return slides

def add_title_slide(prs, title_text, subtitle_text=''):
    # 使用 layout 0 (Title Slide)
    layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(layout)
    try:
        title = slide.shapes.title
        title.text = title_text
        title.text_frame.paragraphs[0].font.name = TITLE_FONT
        title.text_frame.paragraphs[0].font.size = Pt(40)
        title.text_frame.paragraphs[0].font.color.rgb = TITLE_COLOR
    except Exception:
        pass
    try:
        subtitle = slide.placeholders[1]
        subtitle.text = subtitle_text
        subtitle.text_frame.paragraphs[0].font.name = BODY_FONT
        subtitle.text_frame.paragraphs[0].font.size = Pt(14)
    except Exception:
        pass

def add_content_slide(prs, title_text, body_text):
    # 使用 layout 1 (Title and Content)
    layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(layout)
    try:
        slide.shapes.title.text = title_text
        slide.shapes.title.text_frame.paragraphs[0].font.name = TITLE_FONT
        slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(28)
        slide.shapes.title.text_frame.paragraphs[0].font.color.rgb = TITLE_COLOR
    except Exception:
        pass

    # 找到第一个空的文本框作为正文
    body_tf = None
    for shape in slide.shapes:
        if shape.has_text_frame and not shape.text_frame.text:
            body_tf = shape.text_frame
            break
    if not body_tf:
        left = Inches(1)
        top = Inches(1.8)
        width = Inches(8)
        height = Inches(4)
        txBox = slide.shapes.add_textbox(left, top, width, height)
        body_tf = txBox.text_frame

    # 处理正文为多段或带 '- ' 的 bullet
    lines = [ln.strip() for ln in body_text.splitlines() if ln.strip()]
    if not lines:
        return
    body_tf.clear()
    for i, ln in enumerate(lines):
        if i == 0:
            p = body_tf.paragraphs[0]
        else:
            p = body_tf.add_paragraph()
        # 如果行以 '- ' 或 '·' 开始，设为项目符号
        if ln.startswith('- ') or ln.startswith('· '):
            p.text = ln[2:].strip()
            p.level = 0
        else:
            p.text = ln
        p.font.name =

